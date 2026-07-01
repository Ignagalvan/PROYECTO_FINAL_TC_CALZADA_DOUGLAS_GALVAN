from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "Informe Tecnico - plantilla.docx"
OUTPUT = ROOT / "informe" / "Informe_Tecnico_Final_CON_MANUAL.docx"


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def add_body(doc, text="", bold_lead=None, space_after=6):
    paragraph = doc.add_paragraph(style="Cuerpo")
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text, page_break=False):
    paragraph = doc.add_paragraph(style="Título")
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.add_run(text)
    set_keep_with_next(paragraph)
    return paragraph


def add_subheading(doc, text):
    paragraph = doc.add_paragraph(style="Cuerpo")
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    set_keep_with_next(paragraph)
    return paragraph


def add_code(doc, lines):
    for line in lines:
        paragraph = doc.add_paragraph(style="Cuerpo")
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)


doc = Document(TEMPLATE)

# Conserva la portada y el encabezado de resumen de la plantilla.
for paragraph in list(doc.paragraphs)[30:]:
    remove_paragraph(paragraph)

summary_heading = doc.paragraphs[29]
summary_heading.paragraph_format.page_break_before = True
summary_heading.paragraph_format.space_after = Pt(8)
set_keep_with_next(summary_heading)

add_body(
    doc,
    "El presente informe describe el diseño y la implementación de un compilador "
    "educativo para un subconjunto de C++. El sistema fue desarrollado en Java "
    "con ANTLR4 y organiza el procesamiento en análisis léxico, sintáctico y "
    "semántico, generación de código intermedio y optimización.",
)
add_body(
    doc,
    "La solución reconoce variables, arreglos básicos, expresiones, estructuras "
    "de control y funciones. Además, administra ámbitos mediante una tabla de "
    "símbolos, genera instrucciones de tres direcciones y guarda el resultado "
    "antes y después de la optimización. La verificación final ejecutó 17 pruebas "
    "automatizadas sin fallos ni errores.",
)

add_heading(doc, "1. INTRODUCCIÓN")
add_body(
    doc,
    "El objetivo del proyecto fue aplicar de forma práctica los conceptos "
    "centrales de Técnicas de Compilación. Debido a que C++ completo posee una "
    "gramática y una semántica demasiado amplias para el alcance del trabajo, se "
    "definió un lenguaje reducido que conserva construcciones suficientes para "
    "estudiar las distintas fases de compilación.",
)
add_body(
    doc,
    "El compilador recibe un archivo .txt o .cpp y procesa cada etapa en orden. "
    "Si encuentra un error crítico, detiene el flujo y muestra la línea, la "
    "columna y una descripción. Cuando la entrada es válida, presenta los tokens, "
    "el árbol de parseo y la tabla de símbolos, y genera archivos con el código "
    "intermedio y optimizado.",
)

add_subheading(doc, "1.1 Objetivo general")
add_body(
    doc,
    "Diseñar un compilador modular capaz de reconocer, validar y transformar "
    "programas escritos en un subconjunto de C++, mostrando de manera clara las "
    "fases internas y sus resultados.",
)

add_subheading(doc, "1.2 Objetivos específicos")
objectives = [
    ("Gramática: ", "definir tokens, sentencias, funciones y expresiones con precedencia."),
    ("Validación: ", "detectar errores léxicos, sintácticos y semánticos."),
    ("Símbolos: ", "administrar variables, funciones y ámbitos anidados."),
    ("Traducción: ", "generar código de tres direcciones con temporales y etiquetas."),
    ("Optimización: ", "aplicar transformaciones simples sobre el código intermedio."),
    ("Pruebas: ", "verificar los componentes principales mediante JUnit y Maven."),
]
for lead, text in objectives:
    add_body(doc, text, bold_lead=lead, space_after=3)

add_heading(doc, "2. DESARROLLO", page_break=True)

add_subheading(doc, "2.1 Alcance del lenguaje")
add_body(
    doc,
    "La gramática reconoce los tipos int, float, double, char, string, bool y "
    "void. Permite declaraciones e inicializaciones opcionales, asignaciones, "
    "arreglos unidimensionales, cout, expresiones aritméticas, relacionales y "
    "lógicas, if-else, while, for, break, continue, funciones, parámetros, "
    "llamadas y return.",
)
add_body(
    doc,
    "El sistema no intenta ser compatible con todo C++. No incluye clases, "
    "punteros, referencias, plantillas ni generación de código máquina. Su "
    "finalidad es académica: mostrar un pipeline de compilación comprensible y "
    "extensible.",
)

add_subheading(doc, "2.2 Arquitectura general")
add_body(
    doc,
    "La clase App funciona como coordinadora. Carga el archivo, ejecuta cada fase "
    "y evita continuar cuando existen errores. ANTLR4 genera MiLenguajeLexer, "
    "MiLenguajeParser y las clases base del patrón Visitor a partir de "
    "MiLenguaje.g4. Las operaciones propias del proyecto se mantienen en clases "
    "separadas para no modificar el código generado.",
)
add_code(
    doc,
    [
        "Fuente -> Lexer -> Tokens -> Parser -> Árbol de parseo",
        "       -> Visitor semántico -> Tabla de símbolos / errores",
        "       -> CodigoVisitor -> Código de tres direcciones",
        "       -> Optimizador -> Código optimizado / archivos de salida",
    ],
)

add_subheading(doc, "2.3 Análisis léxico")
add_body(
    doc,
    "El lexer transforma los caracteres en tokens. Las palabras reservadas se "
    "declaran antes que ID para impedir que términos como int, while o return se "
    "clasifiquen como identificadores. Los espacios y comentarios se descartan, "
    "mientras que cada token conserva su lexema, línea y columna.",
)
add_body(
    doc,
    "App reemplaza los listeners predeterminados de ANTLR y muestra una tabla de "
    "tokens. La gramática posee una regla OTRO que captura caracteres no "
    "reconocidos; como estos tokens no forman parte de las reglas del parser, "
    "normalmente terminan reportados durante el análisis sintáctico.",
)

add_subheading(doc, "2.4 Análisis sintáctico")
add_body(
    doc,
    "El parser parte de la regla programa, compuesta por cero o más funciones o "
    "sentencias y el marcador EOF. Las expresiones utilizan alternativas "
    "etiquetadas para respetar la precedencia de operadores unarios, "
    "multiplicativos, aditivos, relacionales, de igualdad, AND y OR.",
)
add_code(
    doc,
    [
        "programa : elemento* EOF ;",
        "elemento : funcion | sentencia ;",
        "funcion  : tipo ID '(' parametros? ')' bloque ;",
        "sentencia: declaracion | asignacion | if | while | for | return | bloque ;",
    ],
)
add_body(
    doc,
    "Si la estructura es válida, ANTLR produce un árbol de parseo concreto. Al "
    "final de una compilación exitosa, TreeViewer abre una ventana Swing para "
    "visualizarlo.",
)

add_subheading(doc, "2.5 Análisis semántico")
add_body(
    doc,
    "SemanticAnalyzerVisitor recorre el árbol y utiliza SymbolTable para "
    "registrar variables, parámetros y funciones. Cada Scope mantiene sus "
    "símbolos y una referencia al ámbito padre. La resolución comienza en el "
    "ámbito actual y continúa hacia el exterior, por lo que se admiten variables "
    "locales sin perder el acceso a símbolos globales.",
)
add_body(
    doc,
    "La fase detecta declaraciones duplicadas, variables o funciones no "
    "declaradas, parámetros repetidos y usos de break, continue o return fuera "
    "de contexto. Las funciones se registran antes de recorrer los cuerpos, lo "
    "que permite realizar llamadas antes de su definición textual.",
)
add_body(
    doc,
    "La implementación todavía no compara tipos de expresiones, asignaciones, "
    "argumentos y retornos. Tampoco valida la cantidad de argumentos contra la "
    "firma de una función. La infraestructura admite warnings, pero actualmente "
    "no existen reglas semánticas que los produzcan.",
)

add_subheading(doc, "2.6 Tabla de símbolos")
add_body(
    doc,
    "Symbol representa una variable o función mediante nombre, tipo, categoría "
    "y nivel de ámbito. SymbolTable ofrece operaciones de declaración, búsqueda "
    "y entrada o salida de ámbitos. La tabla que se imprime al finalizar muestra "
    "los símbolos conservados en el ámbito global; los símbolos locales se usan "
    "durante el recorrido, pero no se guardan en un historial para mostrarlos.",
)

add_subheading(doc, "2.7 Generación de código intermedio")
add_body(
    doc,
    "CodigoVisitor convierte las construcciones del árbol en instrucciones de "
    "tres direcciones. GeneradorCodigo crea temporales t0, t1, etc. para "
    "resultados parciales y etiquetas L0, L1, etc. para saltos. Las pilas de "
    "etiquetas de break y continue permiten traducir bucles anidados.",
)
add_code(
    doc,
    [
        "t0 = numeros[0] + numeros[1]",
        "temp = t0",
        "param temp",
        "param 5",
        "t1 = call sumar",
        "if !t2 goto L0",
    ],
)
add_body(
    doc,
    "Una limitación concreta es que visitDeclaracion emite la instrucción declare, "
    "pero no genera la asignación cuando la declaración posee inicializador. Por "
    "ejemplo, int x = 5 produce la declaración de x, pero no la instrucción x = 5.",
)

add_subheading(doc, "2.8 Optimización")
add_body(
    doc,
    "Las estrategias implementan la interfaz OptimizadorIntermedio. El proyecto "
    "incluye eliminación de código inalcanzable después de goto o return, "
    "propagación de constantes, simplificación de expresiones y eliminación de "
    "autoasignaciones. App utiliza por defecto OptimizadorCodigoMuerto.",
)
add_body(
    doc,
    "La selección se realiza manualmente en App.java dejando activa una sola "
    "implementación. Por lo tanto, todavía no existe una cadena configurable que "
    "aplique varias optimizaciones en una misma compilación.",
)

add_subheading(doc, "2.9 Salidas y diagnósticos")
add_body(
    doc,
    "SalidaCompilador centraliza los mensajes importantes. Utiliza verde para "
    "[EXITO], amarillo para [WARNING] y rojo para [ERROR]. Esta organización "
    "mejora la lectura durante pruebas y exposiciones, aunque no modifica el "
    "resultado semántico del compilador.",
)
add_body(
    doc,
    "El código intermedio se guarda con el sufijo _codigo_intermedio.txt y el "
    "resultado de la estrategia seleccionada con _codigo_optimizado.txt. La "
    "consola también presenta el número de tokens y un resumen de errores por fase.",
)

add_subheading(doc, "2.10 Ejemplo representativo")
add_body(
    doc,
    "El archivo EJ_FINAL.txt incluye variables globales, una función con "
    "parámetros, arreglos, operaciones aritméticas, una llamada con retorno y una "
    "estructura if. Este caso permite comprobar en una sola ejecución la mayor "
    "parte del recorrido implementado.",
)
add_code(
    doc,
    [
        "int sumar(int a, int b) {",
        "    int resultado;",
        "    resultado = a + b;",
        "    return resultado;",
        "}",
        "estado = sumar(temp, 5);",
    ],
)

add_subheading(doc, "2.11 Pruebas y resultados")
add_body(
    doc,
    "La suite se ejecutó con mvn test el 12 de junio de 2026. El resultado fue "
    "BUILD SUCCESS con 17 pruebas, 0 fallos, 0 errores y 0 omitidas. La cantidad "
    "se distribuye en AppTest (1), SalidaCompiladorTest (3), "
    "SemanticAnalyzerVisitorTest (6) y SymbolTableTest (7).",
)
add_body(
    doc,
    "Las pruebas cubren la declaración y resolución de símbolos, duplicados, "
    "ámbitos anidados, funciones, parámetros, errores semánticos y colores ANSI. "
    "Aún faltan pruebas específicas para CodigoVisitor y para cada estrategia de "
    "optimización.",
)

add_subheading(doc, "2.12 Limitaciones y mejoras futuras")
limitations = [
    ("AST: ", "el sistema trabaja directamente sobre el árbol de parseo de ANTLR."),
    ("Tipos: ", "no existe inferencia ni verificación completa de compatibilidad."),
    ("Funciones: ", "no se validan cantidad y tipo de argumentos ni retorno."),
    ("Arreglos: ", "no se controla el rango ni el tipo del índice."),
    ("Inicialización: ", "el generador omite la asignación inicial de declaraciones."),
    ("Ámbitos: ", "la impresión final solo conserva los símbolos globales."),
    ("Optimización: ", "solo puede seleccionarse una estrategia por ejecución."),
    ("Backend: ", "no se genera código máquina, bytecode ni ensamblador."),
]
for lead, text in limitations:
    add_body(doc, text, bold_lead=lead, space_after=3)

add_heading(doc, "3. CONCLUSIONES", page_break=True)
add_body(
    doc,
    "El proyecto integra correctamente las etapas principales de un compilador "
    "educativo. ANTLR4 redujo la complejidad de implementar el lexer y el parser, "
    "mientras que el patrón Visitor permitió separar la semántica y la generación "
    "de código de las clases generadas automáticamente.",
)
add_body(
    doc,
    "La solución procesa un subconjunto significativo de C++, administra ámbitos "
    "y produce una representación intermedia comprensible. La división en "
    "componentes facilita identificar errores y permite extender el sistema sin "
    "modificar todo el flujo.",
)
add_body(
    doc,
    "Los objetivos principales fueron alcanzados y las 17 pruebas confirman el "
    "funcionamiento de los componentes evaluados. Para una versión más completa "
    "sería necesario incorporar un AST propio, un sistema de tipos, validación de "
    "firmas, pruebas de generación y un pipeline de optimización configurable.",
)

add_heading(doc, "4. BIBLIOGRAFÍA")
references = [
    "Aho, A. V., Lam, M. S., Sethi, R. y Ullman, J. D. (2007). Compilers: Principles, Techniques, and Tools. Pearson.",
    "Parr, T. (2013). The Definitive ANTLR 4 Reference. Pragmatic Bookshelf.",
    "ANTLR. Documentación oficial. https://www.antlr.org/",
    "Oracle. Documentación de Java. https://docs.oracle.com/en/java/",
    "Apache Maven Project. Documentación oficial. https://maven.apache.org/",
]
for reference in references:
    add_body(doc, reference, space_after=3)

add_heading(doc, "5. FUENTES DE DATOS")
add_body(
    doc,
    "La información técnica fue contrastada con el código fuente del proyecto: "
    "MiLenguaje.g4, App.java, SemanticAnalyzerVisitor, SymbolTable, Scope, "
    "CodigoVisitor, GeneradorCodigo, las clases de optimización y las pruebas JUnit.",
)
add_body(
    doc,
    "Repositorio: https://github.com/Ignagalvan/"
    "PROYECTO_FINAL_TC_CALZADA_DOUGLAS_GALVAN",
)
add_body(
    doc,
    "Verificación local: mvn test ejecutado el 12 de junio de 2026. Resultado: "
    "17 pruebas, 0 fallos, 0 errores y 0 omitidas.",
)

add_heading(doc, "6. GLOSARIO", page_break=True)
terms = [
    ("ANTLR4: ", "herramienta que genera analizadores a partir de una gramática."),
    ("Lexer: ", "componente que agrupa los caracteres de entrada en tokens."),
    ("Parser: ", "componente que valida la estructura definida por la gramática."),
    ("Token: ", "unidad léxica como una palabra reservada, operador o literal."),
    ("Parse tree: ", "árbol concreto construido según las reglas reconocidas."),
    ("AST: ", "representación abstracta del programa, no implementada actualmente."),
    ("Visitor: ", "patrón utilizado para recorrer el árbol y ejecutar una operación."),
    ("Scope: ", "ámbito que contiene símbolos y referencia a un ámbito exterior."),
    ("Tabla de símbolos: ", "estructura que registra variables y funciones."),
    ("Código de tres direcciones: ", "representación intermedia con operaciones simples."),
    ("Temporal: ", "variable auxiliar usada para resultados parciales."),
    ("Optimización: ", "transformación que mejora el código sin alterar su resultado."),
]
for lead, description in terms:
    add_body(doc, description, bold_lead=lead, space_after=3)

add_heading(doc, "7. MANUAL DE USUARIO", page_break=True)

add_subheading(doc, "A.1 Requisitos")
requirements = [
    ("Java: ", "JDK 8 o una versión compatible."),
    ("Maven: ", "necesario para compilar, ejecutar las pruebas y generar el JAR."),
    ("Entrada: ", "archivo de código fuente con extensión .txt o .cpp."),
    ("Interfaz gráfica: ", "necesaria únicamente para visualizar el árbol con TreeViewer."),
]
for lead, description in requirements:
    add_body(doc, description, bold_lead=lead, space_after=3)

add_subheading(doc, "A.2 Compilación del proyecto")
add_body(
    doc,
    "Desde una terminal, ubicarse en la carpeta demo y ejecutar Maven. Este "
    "proceso genera las clases de ANTLR, compila Java, ejecuta las pruebas y crea "
    "el JAR con sus dependencias.",
)
add_code(
    doc,
    [
        "cd demo",
        "mvn clean package",
    ],
)

add_subheading(doc, "A.3 Ejecución")
add_body(
    doc,
    "El programa requiere exactamente una ruta de archivo como argumento. Puede "
    "utilizarse EJ_FINAL.txt para una ejecución válida o EJ_ERROR.txt para "
    "observar el reporte de errores.",
)
add_code(
    doc,
    [
        "java -jar target/demo-1.0-jar-with-dependencies.jar EJ_FINAL.txt",
        "java -jar target/demo-1.0-jar-with-dependencies.jar EJ_ERROR.txt",
    ],
)

add_subheading(doc, "A.4 Interpretación de la ejecución")
add_body(
    doc,
    "Primero se muestra la tabla de tokens con tipo, lexema, línea y columna. Si "
    "no existen errores, continúa el análisis sintáctico y luego el semántico. "
    "Después se imprime la tabla de símbolos global, el código intermedio, el "
    "resumen de optimización y el resultado final.",
)
messages = [
    ("[EXITO]: ", "la fase terminó correctamente y el proceso puede continuar."),
    ("[WARNING]: ", "situación no crítica; actualmente no hay reglas que generen advertencias semánticas."),
    ("[ERROR]: ", "la compilación se detiene hasta corregir la causa informada."),
]
for lead, description in messages:
    add_body(doc, description, bold_lead=lead, space_after=3)

add_subheading(doc, "A.5 Archivos generados")
add_body(
    doc,
    "Para una entrada llamada programa.txt, la aplicación crea los siguientes "
    "archivos en la misma ubicación:",
)
files = [
    ("programa_codigo_intermedio.txt: ", "instrucciones generadas antes de optimizar."),
    ("programa_codigo_optimizado.txt: ", "resultado de la estrategia de optimización activa."),
]
for lead, description in files:
    add_body(doc, description, bold_lead=lead, space_after=3)

add_subheading(doc, "A.6 Selección de la optimización")
add_body(
    doc,
    "La versión actual no permite elegir la estrategia desde la línea de comandos. "
    "La selección se realiza en App.java instanciando una implementación de "
    "OptimizadorIntermedio. De forma predeterminada se utiliza "
    "OptimizadorCodigoMuerto; solamente una opción debe quedar activa.",
)

add_subheading(doc, "A.7 Problemas frecuentes")
problems = [
    ("No se encuentra el archivo: ", "comprobar la ruta y el nombre enviados al JAR."),
    ("Error sintáctico: ", "revisar punto y coma, paréntesis, llaves y la posición indicada."),
    ("Variable no declarada: ", "declararla antes de usarla y dentro de un ámbito visible."),
    ("Función no declarada: ", "verificar el nombre y que exista una definición en el programa."),
    ("No aparece el árbol: ", "ejecutar en un entorno de escritorio con soporte gráfico."),
    ("No se ven los colores: ", "utilizar una terminal compatible con secuencias ANSI."),
]
for lead, description in problems:
    add_body(doc, description, bold_lead=lead, space_after=3)

properties = doc.core_properties
properties.title = "Informe técnico - Compilador educativo con ANTLR4"
properties.subject = "Proyecto final de Técnicas de Compilación"
properties.author = "Tomás Calzada, Octavio Douglas e Ignacio Galvan"
properties.keywords = "ANTLR4, Java, compilador, semántica, código intermedio, optimización"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
